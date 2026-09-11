"""PyQt6 GUI로 네이버 금융 코스피200 편입종목을 표시한다."""

from urllib.parse import urljoin

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import requests
from bs4 import BeautifulSoup
from PyQt6.QtCore import QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


INDEX_URL = "https://finance.naver.com/sise/sise_index.naver?code=KPI200"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 Chrome/131.0 Safari/537.36"
    )
}
EXPECTED_COLUMNS = [
    "종목별",
    "현재가",
    "전일비",
    "등락률",
    "거래량",
    "거래대금(백만)",
    "시가총액(억)",
]
OUTPUT_FILE = "naverResult.xlsx"


def get_soup(url):
    """URL의 HTML을 받아 BeautifulSoup 객체로 반환한다."""
    response = requests.get(url, headers=HEADERS, timeout=15)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or response.encoding
    return BeautifulSoup(response.text, "html.parser")


def clean_text(node):
    """태그 안의 줄바꿈과 연속 공백을 정리한다."""
    return " ".join(node.get_text(" ", strip=True).split())


def find_entry_table(soup, page=1):
    """지정한 페이지의 편입종목 iframe 안에 있는 표를 찾는다."""
    iframe = soup.select_one(
        'iframe[title*="편입종목"], iframe[src*="entryJongmok"]'
    )
    if iframe is None or not iframe.get("src"):
        raise RuntimeError("편입종목상위 iframe을 찾지 못했습니다.")

    iframe_url = urljoin(INDEX_URL, iframe["src"])
    if page > 1:
        iframe_url = f"{iframe_url}&page={page}"
    iframe_soup = get_soup(iframe_url)
    table = iframe_soup.select_one("table.type_1")
    if table is None:
        raise RuntimeError("편입종목상위 표를 찾지 못했습니다.")
    return table


def crawl_kospi200(limit=None):
    """모든 페이지의 편입종목상위 정보를 리스트로 반환한다."""
    index_soup = get_soup(INDEX_URL)
    stocks = []
    columns = EXPECTED_COLUMNS

    for page in range(1, 22):
        table = find_entry_table(index_soup, page)
        rows = table.find_all("tr")
        if not rows:
            break

        header_cells = rows[0].find_all("th")
        page_columns = [clean_text(cell).replace(" ", "") for cell in header_cells]
        if page_columns != columns:
            raise RuntimeError(f"예상하지 못한 표 컬럼입니다: {page_columns}")

        page_stock_count = 0
        for row in rows[1:]:
            cells = row.find_all("td")
            if len(cells) != len(columns):
                continue

            name_link = cells[0].select_one("a")
            if name_link is None:
                continue

            stock = {
                column: clean_text(cell)
                for column, cell in zip(columns, cells)
            }
            stock["종목코드"] = name_link.get("href", "").split("code=")[-1]
            stock["종목URL"] = urljoin(INDEX_URL, name_link.get("href", ""))
            stocks.append(stock)
            page_stock_count += 1

            if limit is not None and len(stocks) >= limit:
                return stocks[:limit]

        if page_stock_count == 0:
            break

    return stocks


def save_stocks_to_excel(stocks, file_path=OUTPUT_FILE):
    """크롤링 결과를 Excel 파일로 저장한다."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "KOSPI200"
    worksheet.append(EXPECTED_COLUMNS)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")

    for stock in stocks:
        worksheet.append([stock[column] for column in EXPECTED_COLUMNS])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    column_widths = [18, 14, 14, 12, 16, 18, 18]
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width

    workbook.save(file_path)


class CrawlThread(QThread):
    """네이버 금융 요청을 GUI와 별도의 스레드에서 실행한다."""

    completed = pyqtSignal(list)
    failed = pyqtSignal(str)

    def run(self):
        try:
            stocks = crawl_kospi200()
            save_stocks_to_excel(stocks)
            self.completed.emit(stocks)
        except requests.RequestException as error:
            self.failed.emit(f"네이버 금융 페이지 요청에 실패했습니다.\n{error}")
        except Exception as error:
            self.failed.emit(f"크롤링 중 오류가 발생했습니다.\n{error}")


class Kospi200Window(QMainWindow):
    """코스피200 편입종목을 QTableWidget으로 보여주는 창."""

    def __init__(self):
        super().__init__()
        self.worker = None
        self.setWindowTitle("코스피200 편입종목상위")
        self.resize(1100, 650)
        self.setup_ui()
        self.load_stocks()

    def setup_ui(self):
        self.status_label = QLabel("데이터를 불러오는 중입니다...")
        self.refresh_button = QPushButton("새로고침")
        self.refresh_button.clicked.connect(self.load_stocks)

        self.table = QTableWidget(0, len(EXPECTED_COLUMNS))
        self.table.setHorizontalHeaderLabels(EXPECTED_COLUMNS)
        self.table.setAlternatingRowColors(True)
        self.table.setSortingEnabled(True)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.resizeColumnsToContents()

        layout = QVBoxLayout()
        layout.addWidget(self.status_label)
        layout.addWidget(self.refresh_button)
        layout.addWidget(self.table)

        central_widget = QWidget()
        central_widget.setLayout(layout)
        self.setCentralWidget(central_widget)

    def load_stocks(self):
        """전체 편입종목을 요청하고 표에 표시한다."""
        if self.worker is not None and self.worker.isRunning():
            return

        self.refresh_button.setEnabled(False)
        self.status_label.setText("데이터를 불러오는 중입니다...")
        self.worker = CrawlThread()
        self.worker.completed.connect(self.display_stocks)
        self.worker.failed.connect(self.show_error)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.start()

    def display_stocks(self, stocks):
        """크롤링 결과를 QTableWidget에 채운다."""
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(stocks))

        for row_index, stock in enumerate(stocks):
            for column_index, column in enumerate(EXPECTED_COLUMNS):
                self.table.setItem(
                    row_index,
                    column_index,
                    QTableWidgetItem(stock[column]),
                )

        self.table.resizeColumnsToContents()
        self.table.setSortingEnabled(True)
        self.status_label.setText(f"총 {len(stocks)}개 종목을 불러왔습니다.")
        self.refresh_button.setEnabled(True)

    def show_error(self, message):
        """크롤링 오류를 상태 표시와 메시지 상자로 알린다."""
        self.status_label.setText("데이터를 불러오지 못했습니다.")
        self.refresh_button.setEnabled(True)
        QMessageBox.critical(self, "오류", message)


def main():
    app = QApplication([])
    window = Kospi200Window()
    window.show()
    app.exec()


if __name__ == "__main__":
    main()