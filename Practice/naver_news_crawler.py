"""PyQt6 기반 네이버 뉴스 크롤러."""

import sys
from datetime import datetime
from urllib.parse import urljoin

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PyQt6.QtCore import QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
    QCheckBox,
)
import requests
from bs4 import BeautifulSoup


NAVER_SEARCH_URL = "https://search.naver.com/search.naver"
DEFAULT_QUERY = "반도체"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 Chrome/131.0 Safari/537.36"
    )
}


def get_soup(url, params=None):
    """URL의 HTML을 받아 BeautifulSoup 객체로 반환한다."""
    response = requests.get(
        url,
        params=params,
        headers=HEADERS,
        timeout=15,
    )
    response.raise_for_status()
    response.encoding = response.apparent_encoding or response.encoding
    return BeautifulSoup(response.text, "html.parser")


def clean_text(node):
    """태그 안의 줄바꿈과 공백을 정리한다."""
    if node is None:
        return ""
    return " ".join(node.get_text(" ", strip=True).split())


def remove_new_window_text(text):
    return text.replace("새 창 열림", "").strip()


def get_area_link(container, area_name):
    """네이버의 data-nlog-area 값으로 카드 내부 링크를 찾는다."""
    return container.select_one(f'a[data-nlog-area$=".{area_name}"]')


def extract_article_body(url):
    """언론사 원문에서 대표적인 기사 본문 영역을 추출한다."""
    try:
        soup = get_soup(url)
    except requests.RequestException as error:
        return f"본문 요청 실패: {error}"

    body_selectors = (
        "#dic_area",
        "article",
        ".newsct_article",
        ".article_body",
        ".article-body",
        ".article_view",
        ".article-view",
    )
    for selector in body_selectors:
        body = soup.select_one(selector)
        text = clean_text(body)
        if text:
            return text

    return "본문 영역을 찾지 못했습니다. 언론사별 선택자를 추가해야 합니다."


def crawl_news(query=DEFAULT_QUERY, limit=10, include_body=False):
    """네이버 뉴스 검색 결과를 수집한다."""
    params = {
        "where": "nexearch",
        "sm": "top_hty",
        "fbm": "0",
        "ie": "utf8",
        "query": query,
    }
    soup = get_soup(NAVER_SEARCH_URL, params=params)
    articles = []
    news_list = soup.select_one(".fds-news-item-list-desk")
    if news_list is None:
        return articles

    for card in news_list.find_all("div", recursive=False):
        title_link = get_area_link(card, "tit")
        summary_link = get_area_link(card, "body")
        if title_link is None:
            continue

        article_url = urljoin(NAVER_SEARCH_URL, title_link.get("href", ""))
        articles.append(
            {
                "title": remove_new_window_text(clean_text(title_link)),
                "url": article_url,
                "source": remove_new_window_text(
                    clean_text(card.select_one(".sds-comps-profile-info-title-text"))
                ),
                "date": clean_text(
                    card.select_one(".sds-comps-profile-info-subtext")
                ),
                "summary": remove_new_window_text(clean_text(summary_link)),
            }
        )
        if len(articles) >= limit:
            break

    for article in articles:
        if include_body:
            article["content"] = extract_article_body(article["url"])

    return articles


def save_articles_to_excel(articles, file_path, query):
    """크롤링한 기사를 Excel 파일로 저장한다."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Naver News"

    headers = ["제목", "원문 URL", "언론사", "작성 시각", "요약", "본문"]
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor="2F5597")

    for article in articles:
        worksheet.append(
            [
                article.get("title", ""),
                article.get("url", ""),
                article.get("source", ""),
                article.get("date", ""),
                article.get("summary", ""),
                article.get("content", ""),
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[1].hyperlink = row[1].value
        row[1].style = "Hyperlink"

    column_widths = {
        "A": 42,
        "B": 65,
        "C": 16,
        "D": 14,
        "E": 70,
        "F": 100,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.oddFooter.center.text = f"검색어: {query}"
    worksheet.oddFooter.right.text = datetime.now().strftime("%Y-%m-%d %H:%M")
    workbook.save(file_path)


class CrawlThread(QThread):
    """크롤링과 Excel 저장을 백그라운드에서 실행한다."""

    completed = pyqtSignal(list, str)
    failed = pyqtSignal(str)

    def __init__(self, query, limit, include_body, output_path):
        super().__init__()
        self.query = query
        self.limit = limit
        self.include_body = include_body
        self.output_path = output_path

    def run(self):
        try:
            articles = crawl_news(self.query, self.limit, self.include_body)
            save_articles_to_excel(articles, self.output_path, self.query)
            self.completed.emit(articles, self.output_path)
        except requests.RequestException as error:
            self.failed.emit(f"네이버 페이지 요청에 실패했습니다.\n{error}")
        except Exception as error:
            self.failed.emit(f"처리 중 오류가 발생했습니다.\n{error}")


class NewsCrawlerWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.worker = None
        self.setWindowTitle("네이버 뉴스 크롤러")
        self.resize(1100, 700)
        self.setup_ui()

    def setup_ui(self):
        query_input = QLineEdit(DEFAULT_QUERY)
        limit_input = QSpinBox()
        limit_input.setRange(1, 100)
        limit_input.setValue(10)
        body_check = QCheckBox("기사 본문도 수집")
        output_input = QLineEdit("naverResult.xlsx")
        browse_button = QPushButton("찾아보기")
        crawl_button = QPushButton("뉴스 크롤링 시작")
        status_label = QLabel("검색어와 저장 경로를 설정한 뒤 시작하세요.")

        form = QFormLayout()
        form.addRow("검색어", query_input)
        form.addRow("기사 수", limit_input)
        form.addRow("옵션", body_check)

        output_layout = QHBoxLayout()
        output_layout.addWidget(output_input)
        output_layout.addWidget(browse_button)
        form.addRow("Excel 저장 경로", output_layout)

        table = QTableWidget(0, 5)
        table.setHorizontalHeaderLabels(["제목", "언론사", "작성 시각", "요약", "원문 URL"])
        table.setAlternatingRowColors(True)
        table.setWordWrap(True)
        table.horizontalHeader().setStretchLastSection(True)
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)

        layout = QVBoxLayout()
        layout.addLayout(form)
        layout.addWidget(crawl_button)
        layout.addWidget(status_label)
        layout.addWidget(table)

        central_widget = QWidget()
        central_widget.setLayout(layout)
        self.setCentralWidget(central_widget)

        browse_button.clicked.connect(
            lambda: self.select_output_path(output_input)
        )
        crawl_button.clicked.connect(
            lambda: self.start_crawling(
                query_input, limit_input, body_check, output_input,
                crawl_button, status_label, table,
            )
        )

    @staticmethod
    def select_output_path(output_input):
        file_path, _ = QFileDialog.getSaveFileName(
            None,
            "Excel 파일 저장",
            output_input.text() or "naverResult.xlsx",
            "Excel 파일 (*.xlsx)",
        )
        if file_path:
            output_input.setText(file_path)

    def start_crawling(
        self, query_input, limit_input, body_check, output_input,
        crawl_button, status_label, table,
    ):
        query = query_input.text().strip()
        output_path = output_input.text().strip()
        if not query:
            QMessageBox.warning(self, "입력 확인", "검색어를 입력하세요.")
            return
        if not output_path:
            QMessageBox.warning(self, "입력 확인", "Excel 저장 경로를 입력하세요.")
            return

        table.setRowCount(0)
        crawl_button.setEnabled(False)
        status_label.setText("크롤링 중입니다. 잠시 기다려 주세요.")
        self.worker = CrawlThread(
            query, limit_input.value(), body_check.isChecked(), output_path
        )
        self.worker.completed.connect(
            lambda articles, path: self.show_results(
                articles, path, crawl_button, status_label, table
            )
        )
        self.worker.failed.connect(
            lambda message: self.show_error(message, crawl_button, status_label)
        )
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.start()

    @staticmethod
    def show_results(articles, path, crawl_button, status_label, table):
        for article in articles:
            row = table.rowCount()
            table.insertRow(row)
            values = [
                article.get("title", ""),
                article.get("source", ""),
                article.get("date", ""),
                article.get("summary", ""),
                article.get("url", ""),
            ]
            for column, value in enumerate(values):
                table.setItem(row, column, QTableWidgetItem(value))
        table.resizeColumnsToContents()
        crawl_button.setEnabled(True)
        status_label.setText(f"{len(articles)}개 기사를 저장했습니다: {path}")

    @staticmethod
    def show_error(message, crawl_button, status_label):
        crawl_button.setEnabled(True)
        status_label.setText("오류가 발생했습니다.")
        QMessageBox.critical(None, "크롤링 오류", message)


def main():
    app = QApplication(sys.argv)
    window = NewsCrawlerWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()